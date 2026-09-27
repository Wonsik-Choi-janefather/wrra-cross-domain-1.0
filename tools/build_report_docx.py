#!/usr/bin/env python3
"""Build the WRRA Cross Domain 1.0 report from its Markdown source."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "D9D9D9"
TEXT = "202020"


def set_run_font(run, name: str, size: float, bold: bool = False, italic: bool = False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(TEXT)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def set_cell_width(cell, width_inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])
    set_run_font(run, "Aptos", 9)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    subtitle.paragraph_format.space_after = Pt(8)

    for style_name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 12.5, 13, 5),
        ("Heading 3", 11, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in doc.styles:
        fig_style = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        fig_style = doc.styles["Figure Caption"]
    fig_style.font.name = "Aptos"
    fig_style.font.size = Pt(9)
    fig_style.font.italic = True
    fig_style.font.color.rgb = RGBColor.from_string("555555")
    fig_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_style.paragraph_format.space_before = Pt(3)
    fig_style.paragraph_format.space_after = Pt(9)
    fig_style.paragraph_format.keep_with_next = False

    if "Table Caption" not in doc.styles:
        table_style = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = doc.styles["Table Caption"]
    table_style.font.name = "Aptos"
    table_style.font.size = Pt(9)
    table_style.font.bold = True
    table_style.font.color.rgb = RGBColor.from_string(TEXT)
    table_style.paragraph_format.space_before = Pt(8)
    table_style.paragraph_format.space_after = Pt(4)
    table_style.paragraph_format.keep_with_next = True

    if "Code Paragraph" not in doc.styles:
        code_style = doc.styles.add_style("Code Paragraph", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = doc.styles["Code Paragraph"]
    code_style.font.name = "Liberation Mono"
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string(TEXT)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "WRRA Cross Domain 1 0"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, "Aptos", 8.5)
        run.font.color.rgb = RGBColor.from_string("666666")
    add_page_number(section.footer.paragraphs[0])
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""


INLINE_PATTERN = re.compile(r"(\*\*.*?\*\*|`.*?`)")


def add_inline(paragraph, text: str, size: float = 10.5):
    lines = text.replace("  \n", "\n").split("\n")
    for line_index, line in enumerate(lines):
        if line_index:
            paragraph.add_run().add_break()
        position = 0
        for match in INLINE_PATTERN.finditer(line):
            if match.start() > position:
                run = paragraph.add_run(line[position : match.start()])
                set_run_font(run, "Aptos", size)
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                set_run_font(run, "Aptos", size, bold=True)
            else:
                run = paragraph.add_run(token[1:-1])
                set_run_font(run, "Liberation Mono", max(8, size - 1))
            position = match.end()
        if position < len(line):
            run = paragraph.add_run(line[position:])
            set_run_font(run, "Aptos", size)


def add_table(doc: Document, rows: list[list[str]]):
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)

    available = 6.86
    if columns == 3:
        widths = [available * 0.19, available * 0.36, available * 0.45]
    elif columns == 4:
        widths = [available * 0.16, available * 0.28, available * 0.28, available * 0.28]
    elif columns == 5:
        widths = [available * 0.14, available * 0.23, available * 0.19, available * 0.18, available * 0.26]
    else:
        widths = [available / columns] * columns

    font_size = 8.5 if columns <= 5 else 7.7
    for row_index, (row, values) in enumerate(zip(table.rows, rows)):
        for col_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if col_index == 0 or len(value) > 30
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.clear()
            add_inline(paragraph, value, font_size)
            for run in paragraph.runs:
                if row_index == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    table_lines = []
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows, index


def add_image(doc: Document, image_path: Path, caption: str | None, equation: bool):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = bool(caption)
    width = Inches(6.55 if not equation else 5.85)
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=width)
    equation_alt = {
        "equation_common_state.png": "WRRA state and domain transition equations",
        "equation_renderer.png": "Certified Renderer action equivalence equations",
        "equation_grid.png": "DC power flow and line thermal residue equations",
        "equation_scheduling.png": "Sequence dependent setup and machine availability equations",
    }
    alt_text = caption or equation_alt.get(image_path.name, image_path.stem.replace("_", " "))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", alt_text)
    if caption:
        cap = doc.add_paragraph(style="Figure Caption")
        add_inline(cap, caption, 9)
    else:
        paragraph.paragraph_format.space_after = Pt(7)


def build(markdown_path: Path, output_path: Path):
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = (
        "WRRA Cross Domain 1.0 Comparative Exact Validation in Games Power Grids and Scheduling"
    )
    doc.core_properties.author = "Wonsik Choi"
    doc.core_properties.subject = (
        "Exact cross-domain validation of state residue and certified rendering"
    )
    doc.core_properties.keywords = (
        "WRRA, residue, state sufficiency, exact quotient, game, power grid, scheduling"
    )

    index = 0
    cover = True
    paragraph_buffer: list[str] = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = "\n".join(paragraph_buffer).strip()
        paragraph_buffer = []
        if not text:
            return
        if text.startswith("Table "):
            paragraph = doc.add_paragraph(style="Table Caption")
        elif text.startswith("`") and text.endswith("`") and text.count("`") == 2:
            paragraph = doc.add_paragraph(style="Code Paragraph")
        else:
            paragraph = doc.add_paragraph()
        if cover:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(14)
        add_inline(paragraph, text)

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped == "[[PAGE_BREAK]]":
            flush_paragraph()
            doc.add_page_break()
            cover = False
            index += 1
            continue

        image_match = re.fullmatch(r"\[\[(FIGURE|EQUATION):([^|\]]+)(?:\|([^\]]+))?\]\]", stripped)
        if image_match:
            flush_paragraph()
            kind, relative, caption = image_match.groups()
            image_path = (markdown_path.parent / relative).resolve()
            add_image(doc, image_path, caption, kind == "EQUATION")
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if cover:
                if level == 1:
                    paragraph = doc.add_paragraph(style="Title")
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(105)
                    remove_paragraph_border(paragraph)
                    add_inline(paragraph, text, 24)
                    for run in paragraph.runs:
                        set_run_font(run, "Aptos Display", 24, bold=True)
                else:
                    paragraph = doc.add_paragraph(style="Subtitle")
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(paragraph, text, 13 if level == 2 else 11)
            else:
                paragraph = doc.add_paragraph(text, style=f"Heading {level}")
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.*)$", stripped)
        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet_match or number_match:
            flush_paragraph()
            content = (bullet_match or number_match).group(1)
            style = "List Bullet" if bullet_match else "List Number"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, content)
            index += 1
            continue

        paragraph_buffer.append(raw.rstrip())
        index += 1

    flush_paragraph()

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.style.name == "Normal" and paragraph.text:
            paragraph.paragraph_format.widow_control = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
